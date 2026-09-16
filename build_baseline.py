import os
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── helpers ─────────────────────────────────────────────────────────
def fill(h):
    return PatternFill('solid', fgColor=h)

def fnt(bold=False, size=10, color='000000', italic=False):
    return Font(name='Arial', bold=bold, size=size, color=color, italic=italic)

def aln(h='left', wrap=True, v='top'):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def cw(ws, col, w):
    ws.column_dimensions[get_column_letter(col)].width = w

thin  = Side(style='thin',   color='CCCCCC')
bdr   = Border(left=thin, right=thin, top=thin, bottom=thin)

# Renkler
C_HDR    = 'F2F2F2'  # header bg
C_BLUE   = '1F3864'  # koyu lacivert — başlık
C_LBLUE  = 'BDD7EE'  # açık mavi — bölüm başlığı
C_MBLUE  = '2E75B6'  # orta mavi
C_RED    = 'C00000'
C_LRED   = 'FFE0E0'
C_GREEN  = '375623'
C_LGRN   = 'E2EFDA'
C_LYEL   = 'FFFACC'
C_LGRY   = 'F5F5F5'
C_WHITE  = 'FFFFFF'
C_NEW    = 'FFF0CC'   # yeni eklenen satır highlight

HEADERS = [
    'No', 'Bölüm (Section)', 'Gereksinim (Requirement)',
    'Yapılması Gereken Ayar & Değeri (Exact Configuration)',
    'İlgili Dosya / Servis / Komut (Target Component)',
    'Statü (Priority)', 'Güvenlik Gerekçesi & Engellenen Tehdit (Security Rationale)',
    'Değişiklik Durumu (Revision Note)'
]
COL_W = [5, 22, 30, 42, 30, 10, 42, 28]

def make_header_row(ws, row, title, color=C_BLUE):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    c = ws.cell(row=row, column=1, value=title)
    c.font = Font(name='Arial', bold=True, size=13, color=C_WHITE)
    c.fill = fill(color)
    c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)
    ws.row_dimensions[row].height = 28

def make_col_headers(ws, row):
    ws.row_dimensions[row].height = 30
    for ci, h in enumerate(HEADERS, 1):
        c = ws.cell(row=row, column=ci, value=h)
        c.font = fnt(True, 9, C_WHITE)
        c.fill = fill(C_MBLUE)
        c.border = bdr
        c.alignment = aln('center', True, 'center')

def make_section_row(ws, row, section_name):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    c = ws.cell(row=row, column=1, value=f'  {section_name}')
    c.font = fnt(True, 9, C_BLUE)
    c.fill = fill(C_LBLUE)
    c.alignment = aln('left', False, 'center')
    ws.row_dimensions[row].height = 18

def write_row(ws, row, data, is_new=False, is_revised=False):
    ws.row_dimensions[row].height = 52
    row_bg = C_NEW if is_new else (C_LYEL if is_revised else (C_LGRY if row % 2 == 0 else C_WHITE))
    for ci, val in enumerate(data, 1):
        c = ws.cell(row=row, column=ci, value=val)
        c.border = bdr
        c.fill = fill(row_bg)
        if ci == 1:  # No
            c.font = fnt(True, 9, C_BLUE)
            c.alignment = aln('center', False, 'center')
        elif ci == 6:  # Status
            color = C_RED if val and 'Must' in str(val) else C_BLUE
            c.font = fnt(True, 9, color)
            c.alignment = aln('center', False, 'center')
        elif ci == 8:  # Revision
            is_yeni = val and 'YENİ' in str(val).upper()
            c.font = fnt(True, 9, C_RED if is_yeni else C_GREEN)
            c.alignment = aln('left', True, 'top')
        else:
            c.font = fnt(False, 9)
            c.alignment = aln('left', True, 'top')

def setup_columns(ws):
    for i, w in enumerate(COL_W, 1):
        cw(ws, i, w)
    ws.freeze_panes = 'A3'
    ws.sheet_view.showGridLines = False

# ══════════════════════════════════════════════════════════════════
wb = openpyxl.Workbook()
wb.remove(wb.active)

# ══════════════════════════════════════════════════════════════════
# SHEET 1 — LINUX SERVER
# ══════════════════════════════════════════════════════════════════
ws_linux = wb.create_sheet('Linux Server Baseline')
setup_columns(ws_linux)
make_header_row(ws_linux, 1, 'Configuration Baseline Checklist — Linux Server  |  CIS L1 Aligned')
make_col_headers(ws_linux, 2)

linux_data = [
    # (No, Section, Requirement, Exact Config, Target, Status, Rationale, Revision)
    (1,'Installation','Ensure the virtual machine is connected to the appropriate vLAN.',
     'Sunucu ağ arayüzü, şirket ağ topolojisindeki DMZ, Yönetim veya Uygulama vLAN kurallarına uygun subnet IP\'sine atanmalıdır.',
     'Hipervizör (vCenter/ESXi) & Ağ Arayüzü','Must',
     'Farklı güvenlik seviyelerindeki sunucuların mantıksal olarak izole edilmesi ve yetkisiz ağ erişimlerinin engellenmesi.',
     'Mevcut Madde (Onaylandı)'),
    (2,'Access & Authentication','Have a strong password for root account and enforce password complexity.',
     'minlen=21, dcredit=-1, ucredit=-1, lcredit=-1, ocredit=-1, SHA512 algoritması. Root parolası asla e-posta/sohbet üzerinden paylaşılmamalıdır.',
     '/etc/security/pwquality.conf & /etc/login.defs','Must',
     'Kaba kuvvet (brute-force) ve sözlük saldırılarıyla en yetkili işletim sistemi hesabının ele geçirilmesini engeller.',
     'Mevcut Madde (Güçlendirildi: Değerler Eklendi)'),
    (3,'Installation','The name of server must meet relevant naming rules and prevent OS information disclosure.',
     'Sunucu adı kurumsal hostname formatına uygun olmalı. /etc/issue, /etc/issue.net ve /etc/motd dosyalarından OS sürüm bilgileri silinmeli, yasal uyarı banner\'ı eklenmelidir.',
     '/etc/hostname, /etc/issue, /etc/motd','Must',
     'Saldırganların sunucuya bağlanır bağlanmaz işletim sistemi ve çekirdek sürümünü öğrenerek hedefe yönelik exploit aramasını engeller.',
     'Mevcut Madde (Güçlendirildi: Banner/Sürüm Gizleme)'),
    (4,'Installation & Storage','Server should be partitioned correctly; temporary directories must have restrictive mount options.',
     'Uygulamalar ve loglar (/var/log) OS diskinden ayrı olmalı. /tmp, /var/tmp ve /dev/shm bölümleri \'nodev, nosuid, noexec\' bayraklarıyla mount edilmelidir.',
     '/etc/fstab','Must',
     'Saldırganların /tmp veya bellek alanlarına zararlı script/binary yükleyip doğrudan çalıştırmasını (arbitrary code execution) engeller.',
     'Mevcut Madde (Güçlendirildi: nodev,nosuid,noexec Eklendi)'),
    (5,'Network & Security','Host-based firewall has to be enabled and enforcing default deny incoming at all times.',
     'RHEL/Oracle/Rocky/SUSE: systemctl enable --now firewalld\nUbuntu/Debian: ufw enable\nVarsayılan gelen (incoming) tüm trafik DROP/REJECT olmalı, sadece izinli servis portları açılmalıdır.',
     'firewalld / ufw / nftables','Must',
     'Sunucularda yerel güvenlik duvarı ASLA opsiyonel olamaz. Yanal hareket (Lateral Movement) saldırılarını durduran ana kalkandır.',
     'Recommended -> MUST Yapıldı'),
    (6,'Corporate Agents','The machines are secured with the Trend Micro Deep Security / Apex One agent.',
     'ds_agent servisi aktif ve çalışır durumda olmalıdır (systemctl is-active ds_agent). Güncel imza/pattern\'lar otomatik güncellenmelidir.',
     'ds_agent servisi (Trend Micro)','Must',
     'Zararlı yazılım, fidye yazılımı (ransomware) ve yetkisiz işlem girişimlerinin merkezi SOC tarafından anında tespiti ve engellenmesi.',
     'Mevcut Madde (Servis Doğrulaması Eklendi)'),
    (7,'Corporate Agents & Patching','All the packages must be up to date. Keep updated via ManageEngine and official repositories.',
     'uemsagent (ManageEngine) servisi aktif olmalı. RHEL sunucularda subscription-manager status Current olmalı; paket güncellemeleri düzenli uygulanmalıdır.',
     'ManageEngine (uemsagent) & subscription-manager','Must',
     'Bilinen CVE güvenlik zafiyetlerinin hızlıca kapatılması ve sunucunun resmi depolardan dijital imzalı güvenli paket alması.',
     'Mevcut Madde (Servis Doğrulaması Eklendi)'),
    (8,'Access & Authentication','Permissions must be assigned strictly via visudo/sudoers with logging enabled, not uncontrolled group membership.',
     '/etc/sudoers dosyasına "Defaults logfile=\\"/var/log/sudo.log\\"" ve "Defaults requiretty" eklenmeli. Kullanıcılara sınırsız ALL yerine kısıtlı komut yetkisi tanımlanmalıdır.',
     '/etc/sudoers, /etc/sudoers.d/*','Must',
     'Kullanıcıların kontrolsüz root yetkisi almasını önler; çalıştırılan her süper kullanıcı komutunu denetim kaydı altına alır.',
     'Mevcut Madde (Güçlendirildi: Loglama & TTY Eklendi)'),
    (9,'Access & Authentication','Use Active Directory user accounts via SSSD/Realmd for administrative and interactive access.',
     'sssd servisi aktif olmalı (systemctl is-active sssd). Sunucu Active Directory alan adına dahil edilmeli (realm join), erişimler merkezi kurumsal AD gruplarıyla sınırlandırılmalıdır.',
     'sssd servisi & /etc/sssd/sssd.conf','Must',
     'Personel işten ayrıldığında tek noktadan erişim kapatma, parola politikasını merkezileştirme ve merkezi SIEM izlenebilirliği sağlar.',
     'Mevcut Madde (Must yapıldı, SSSD Kriteri Eklendi)'),
    (10,'Access & Authentication','Interactive login must be disabled for local system and service accounts.',
     'UID < 1000 olan tüm sistem/hizmet hesaplarının (daemon, bin, sys, adm, sync vb.) kabukları (shell) /sbin/nologin veya /bin/false olarak ayarlanmalıdır.',
     '/etc/passwd','Must',
     'Saldırganların servis hesaplarını kullanarak sunucu üzerinde interaktif oturum açmasını veya komut çalıştırmasını engeller.',
     'Mevcut Madde (Must yapıldı, Shell Kriteri Eklendi)'),
    (11,'Logging & Auditing','Enable Rsyslog to collect and forward all system logs to the central SIEM / Log server.',
     'rsyslog servisi aktif ve çalışır olmalı. /etc/rsyslog.conf içinde merkezi log sunucusu IP/FQDN hedefi "*.* @log-server.kurum.local:514" tanımlı olmalıdır.',
     'rsyslog servisi & /etc/rsyslog.conf','Must',
     'Sunucu ele geçirilse dahi saldırganın yerel logları silerek izini kaybettirmesini önler; 5651 ve KVKK yasal uyumluluğunu sağlar.',
     'Mevcut Madde (Log İletim Hedefi Eklendi)'),
    (12,'Corporate Agents & Monitoring','All machines are monitored with the SolarWinds agent.',
     'swiagent servisi aktif ve çalışır durumda olmalıdır (systemctl is-active swiagent).',
     'swiagent servisi (SolarWinds Orion)','Must',
     'Kaynak tüketimi (CPU, RAM, Disk doluluğu), kesinti ve yetkisiz servis durdurmalarının anlık olarak SOC/NOC tarafından izlenmesi.',
     'Mevcut Madde (Servis Doğrulaması Eklendi)'),
    (13,'SSH & Remote Access','Direct SSH login for the root account must be strictly prohibited.',
     'PermitRootLogin no\n(Root kullanıcısı doğrudan ağdan SSH ile bağlanamamalı, önce yetkili AD/yerel kullanıcı girmeli, gerekiyorsa sudo kullanılmalıdır.)',
     '/etc/ssh/sshd_config','Must',
     'Kaba kuvvet (brute-force) saldırılarının %99\'u doğrudan "root" hesabını hedefler. Root\'un SSH\'a kapatılması en hayati savunma adımıdır.',
     'YENİ EKLENDİ (Kritik Savunma - CIS 5.2.10)'),
    (14,'SSH & Remote Access','SSH authentication attempts and empty passwords must be restricted.',
     'PermitEmptyPasswords no\nMaxAuthTries 4\nLoginGraceTime 60\nX11Forwarding no',
     '/etc/ssh/sshd_config','Must',
     'Boş parolalı hesapların girişini engeller; başarısız denemeyi 4 ile sınırlayarak parola tahmin saldırılarını durdurur; askıda bağlantıyı 60 sn\'de koparır.',
     'YENİ EKLENDİ (Kritik Savunma - CIS 5.2.5)'),
    (15,'SSH & Remote Access','Inactive terminal and SSH sessions must automatically terminate after 15 minutes.',
     'SSH: ClientAliveInterval 300 & ClientAliveCountMax 3\nShell: /etc/profile.d/timeout.sh içine "export TMOUT=900" ve "readonly TMOUT".',
     '/etc/ssh/sshd_config & /etc/profile.d/timeout.sh','Must',
     'Yöneticinin bilgisayar başından ayrıldığı durumlarda açık unutulan konsol veya SSH oturumlarının ele geçirilmesini (Session Hijacking) engeller.',
     'YENİ EKLENDİ (Kritik Savunma - CIS 5.2.19)'),
    (16,'Network & Kernel (Sysctl)','IP forwarding must be disabled and IP spoofing protections must be enabled.',
     'net.ipv4.ip_forward = 0\nnet.ipv4.conf.all.rp_filter = 1\nnet.ipv4.conf.default.rp_filter = 1\nnet.ipv4.conf.all.accept_redirects = 0',
     '/etc/sysctl.d/99-security.conf','Must',
     'Sunucunun saldırgan tarafından yetkisiz bir router/gateway gibi kullanılarak kurumsal trafiği gizlice iletmesini ve sahte (spoofed) IP paketlerini engeller.',
     'YENİ EKLENDİ (Kritik Savunma - CIS 3.1.1, 3.2.1)'),
    (17,'Network & Kernel (Sysctl)','TCP SYN cookies must be enabled to mitigate Denial of Service (DoS) attacks.',
     'net.ipv4.tcp_syncookies = 1',
     '/etc/sysctl.d/99-security.conf','Must',
     'Sunucu bellek tablosunu sahte bağlantı istekleriyle doldurarak hizmeti durdurmayı hedefleyen SYN Flood DoS saldırılarına karşı çekirdek koruması sağlar.',
     'YENİ EKLENDİ (Kritik Savunma - CIS 3.2.8)'),
    (18,'File System & Integrity','Critical authentication and identity files must have strict ownership and restrictive permissions.',
     '/etc/shadow: chmod 0000 (veya 0640), chown root:root (veya root:shadow)\n/etc/passwd: chmod 0644, chown root:root\n/etc/gshadow: chmod 0000\n/etc/group: chmod 0644',
     '/etc/shadow, /etc/passwd, /etc/gshadow, /etc/group','Must',
     'Parola özetlerinin (hash) yetkisiz kullanıcılar tarafından okunup offline kırılmasını veya değiştirilerek root yetkisi kazanılmasını (privilege escalation) engeller.',
     'YENİ EKLENDİ (Kritik Savunma - CIS 6.1.2 - 6.1.9)'),
    (19,'Logging & Auditing','Kernel audit daemon (auditd) must be active and record critical security events.',
     'auditd servisi aktif olmalı. /etc/shadow değişiklikleri, sistem saati değişiklikleri, kullanıcı ekleme-silme ve yetki yükseltme çağrıları denetlenmelidir.',
     'auditd servisi & /etc/audit/rules.d/audit.rules','Must',
     'Rsyslog sadece çalışan servis logunu tutar; root kullanıcısının sildiği bir dosyayı, saatle oynamasını veya yetkisiz sistem çağrısını ancak auditd yakalar.',
     'YENİ EKLENDİ (Kritik Savunma - CIS 4.1.1 - 4.1.17)'),
    (20,'Access & Authentication','Accounts must be locked after consecutive failed authentication attempts.',
     'pam_faillock (veya pam_tally2): deny=5, unlock_time=900, even_deny_root\n5 hatalı parola denemesinden sonra hesap 15 dakika kilitlenmelidir.',
     '/etc/pam.d/password-auth & /etc/pam.d/system-auth','Must',
     'Sunucu üzerindeki yerel ve servis hesaplarına yönelik otomatik sözlük ve kaba kuvvet parola denemelerini hesabı geçici kilitleyerek engeller.',
     'YENİ EKLENDİ (Kritik Savunma - CIS 5.4.2)'),
    (21,'Services & Daemons','Unencrypted legacy services and unnecessary listening daemons must be removed.',
     'telnet, rsh, rlogin, ypserv (nis), tftp, cups (yazıcı), avahi-daemon (mDNS) servisleri kaldırılmalı veya maskelenmelidir (systemctl mask).',
     'systemd servisleri & paket yöneticisi','Must',
     'Şifresiz açık metin haberleşen protokolleri ve sunucuda gereksiz port açarak saldırı yüzeyi yaratan miras servisleri tamamen yok eder.',
     'YENİ EKLENDİ (Kritik Savunma - CIS 2.1.1 - 2.2.17)'),
    (22,'System Hardening & MAC','Mandatory Access Control (MAC) must be active and enforcing at boot.',
     'Red Hat / Oracle / Rocky: /etc/selinux/config içinde SELINUX=enforcing\nDebian / Ubuntu / SUSE: apparmor.service aktif ve enforce modda olmalıdır.',
     '/etc/selinux/config & apparmor.service','Must',
     'Sıfırıncı gün (0-day) açıklarında bir web sunucusu veya servis hacklense bile saldırganın işletim sisteminin geri kalanına ve diskine erişmesini engeller.',
     'YENİ EKLENDİ (Kritik Savunma - CIS 1.6.1)'),
    (23,'System Hardening & Scheduling','Cron daemon configuration and directories must be restricted to root only.',
     '/etc/crontab ve /etc/cron.* dizinleri: chmod 0700 (veya 0600), chown root:root\n/etc/cron.allow mevcut olmalı, cron.deny kaldırılmalıdır.',
     '/etc/crontab & /etc/cron.*','Must',
     'Linux sistemlerde saldırganların arka kapı (persistence) bırakmak ve root olmak için en çok istismar ettiği zamanlanmış görev manipülasyonunu engeller.',
     'YENİ EKLENDİ (Kritik Savunma - CIS 5.1.1 - 5.1.8)'),
    (24,'System Hardening & Memory','Address Space Layout Randomization (ASLR) enabled and core dumps disabled.',
     'kernel.randomize_va_space = 2\nfs.suid_dumpable = 0\n/etc/security/limits.conf içine "* hard core 0" eklenmelidir.',
     '/etc/sysctl.d/99-security.conf & /etc/security/limits.conf','Must',
     'Bellek taşması (Buffer Overflow) exploitlerini zorlaştırır; çöken servislerin bellek dökümlerinden (core dump) şifre ve sertifika anahtarı sızmasını engeller.',
     'YENİ EKLENDİ (Kritik Savunma - CIS 1.5.1, 1.5.3)'),
    (25,'Access & Authentication','Root must be the only account in the system with UID 0.',
     'awk -F: \'($3 == 0) { print $1 }\' /etc/passwd komutu yalnızca tek bir satır olarak "root" çıktısını vermelidir.',
     '/etc/passwd','Must',
     'Saldırganların yetki yükselttikten sonra sisteme kalıcılık sağlamak için eklediği gizli ikinci root yetkili (UID 0 backdoor) hesapları tespit eder ve engeller.',
     'YENİ EKLENDİ (Kritik Savunma - CIS 5.4.3)'),
]

# Section grupları
LINUX_SECTIONS = {
    1: 'Installation & Network',
    6: 'Corporate Agents & Patching',
    8: 'Access & Authentication',
    11: 'Logging & Auditing',
    13: 'SSH & Remote Access',
    16: 'Network & Kernel Hardening',
    18: 'File System & Integrity',
    19: 'Logging & Auditing (auditd)',
    21: 'Services, Daemons & System Hardening',
}

r = 3
for row_data in linux_data:
    no = row_data[0]
    if no in LINUX_SECTIONS:
        make_section_row(ws_linux, r, LINUX_SECTIONS[no])
        r += 1
    is_new = 'YENİ EKLENDİ' in str(row_data[7])
    is_rev = 'Güçlendirildi' in str(row_data[7]) or '-> MUST' in str(row_data[7])
    write_row(ws_linux, r, list(row_data), is_new=is_new, is_revised=is_rev)
    r += 1

# Renk açıklaması
r += 1
ws_linux.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
c = ws_linux.cell(r, 1, '  Renk Kodu:   Beyaz/Gri = Mevcut madde   |   Sarı = Güçlendirildi / Statü değiştirildi   |   Turuncu = YENİ EKLENDİ')
c.font = fnt(True, 8, '7F3F00')
c.fill = fill('FFF0CC')
c.alignment = aln('left', False, 'center')
ws_linux.row_dimensions[r].height = 18

# ══════════════════════════════════════════════════════════════════
# SHEET 2 — WINDOWS SERVER
# ══════════════════════════════════════════════════════════════════
ws_srv = wb.create_sheet('Windows Server Baseline')
setup_columns(ws_srv)
make_header_row(ws_srv, 1, 'Configuration Baseline Checklist — Windows Server  |  CIS L1 Aligned')
make_col_headers(ws_srv, 2)

server_data = [
    # Mevcut maddeler
    (1,'Installation','Set a complex password for the administrator',
     'En az 21 karakter, büyük/küçük harf, rakam ve özel karakter içermeli. Varsayılan "Administrator" hesabı yeniden adlandırılmalıdır.',
     'Local Security Policy / GPO','Must',
     'Varsayılan admin hesabına yönelik brute-force ve credential stuffing saldırılarını engeller.',
     'Mevcut Madde (Değer Eklendi)'),
    (2,'Installation','After installation, set server roles, apply latest patches, install AV and update signatures.',
     'Windows Update veya WSUS üzerinden tüm güvenlik yamaları uygulanmalı. Trend Micro DS Agent kurulmalı ve güncel tutulmalıdır.',
     'Windows Update / WSUS / ds_agent','Must',
     'Bilinen CVE\'lerin kapatılması ve zararlı yazılımların engellenmesi için temel gereksinim.',
     'Mevcut Madde (Servis Doğrulaması Eklendi)'),
    (3,'Installation','Adding the server to Active Directory is recommended.',
     'Sunucu domain\'e dahil edilmeli. Özel gerekçe yoksa standalone çalıştırılmamalıdır.',
     'Active Directory / Domain','Recommended',
     'Merkezi GPO yönetimi, güvenlik politikası tutarlılığı ve SIEM entegrasyonu için gereklidir.',
     'Mevcut Madde (Onaylandı)'),
    (4,'User Account Security Control','Administrator account shall be renamed.',
     'NewAdministratorName GPO ayarı ile varsayılan "Administrator" hesabı farklı bir isimle değiştirilmeli.',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'Saldırganların ilk denediği hesap adını geçersiz kılar, otomatik saldırıları engeller.',
     'Mevcut Madde (GPO Yolu Eklendi)'),
    (5,'User Account Security Control','Guest account shall be disabled or renamed.',
     'Guest hesabı devre dışı bırakılmalı. Özel ihtiyaç varsa yeniden adlandırılmalıdır.',
     'GPO: Computer Configuration > Windows Settings > Security Settings > Local Policies','Must',
     'Kimliksiz erişim riskini ortadan kaldırır.',
     'Mevcut Madde (Onaylandı)'),
    (6,'Account Security Control','User rights shall meet the principle of least authority.',
     'Her hesabın yalnızca görevi için gereken minimum yetkilere sahip olması sağlanmalıdır. Gereksiz grup üyelikleri kaldırılmalıdır.',
     'GPO: User Rights Assignment','Must',
     'Lateral movement ve privilege escalation saldırılarını sınırlandırır.',
     'Mevcut Madde (Onaylandı)'),
    (7,'Account Security Control','All local accounts\' passwords shall be kept in PAM Tool (LAPS).',
     'Microsoft LAPS kurulmalı ve yapılandırılmalıdır. ADPasswordEncryptionEnabled = 1. PasswordLength minimum 15 karakter. PasswordAgeDays = 30.',
     'Microsoft LAPS / Active Directory','Must',
     'Lokal admin şifresinin tüm sunucularda aynı olmamasını sağlar; credential reuse saldırılarını engeller.',
     'Mevcut Madde (LAPS Detayları Eklendi)'),
    (8,'Password Policy','Enable password policy per "Kullanıcı Kimlik ve Hesap Yönetimi Süreci".',
     'Minimum uzunluk: 21 karakter. Karmaşıklık: Etkin. Maksimum yaş: 90 gün. Geçmiş: 24 şifre. Hesap kilitleme: 5 deneme / 15 dk.',
     'GPO: Default Domain Policy > Password Policy','Must',
     'Zayıf ve tekrar kullanılan şifrelere karşı temel koruma.',
     'Mevcut Madde (Değerler Eklendi)'),
    (9,'Audit Policies','Audit policy change',
     'Success, Failure — GPO üzerinden etkinleştirilmeli.',
     'GPO: Advanced Audit Policy Configuration','Must',
     'Güvenlik politikası değişikliklerinin denetim altında tutulması.',
     'Mevcut Madde (Onaylandı)'),
    (10,'Audit Policies','Audit logon event',
     'Success, Failure — özellikle başarısız giriş denemeleri izlenmelidir.',
     'GPO: Advanced Audit Policy Configuration','Must',
     'Brute-force ve yetkisiz erişim denemelerinin tespiti.',
     'Mevcut Madde (Onaylandı)'),
    (11,'Audit Policies','Audit object access',
     'Failure — kritik dosya ve kayıt defteri erişim hataları loglanmalıdır.',
     'GPO: Advanced Audit Policy Configuration','Must',
     'Yetkisiz dosya/kayıt erişim girişimlerinin tespiti.',
     'Mevcut Madde (Onaylandı)'),
    (12,'Audit Policies','Audit process tracking — Process Creation must include command line.',
     'Success zorunlu. ProcessCreationIncludeCmdLine_Enabled = 1 GPO ile aktif edilmeli. "Optional" statüsü kaldırıldı.',
     'GPO: Administrative Templates > System > Audit Process Creation','Must',
     'Çalıştırılan her komutun tam kaydı olmadan saldırı analizi yapılamaz. SIEM entegrasyonu için kritik.',
     'GÜNCELLEME: Optional -> Must. CIS zorunlu kılıyor.'),
    (13,'Audit Policies','Audit directory service access',
     'Success, Failure — AD nesnelerindeki değişiklikler loglanmalıdır.',
     'GPO: Advanced Audit Policy Configuration','Must',
     'AD üzerindeki yetkisiz değişikliklerin (hesap, grup, GPO) tespiti.',
     'Mevcut Madde (Onaylandı)'),
    (14,'Audit Policies','Audit privilege use',
     'Success, Failure',
     'GPO: Advanced Audit Policy Configuration','Must',
     'Tehlikeli hakların kullanımının izlenmesi.',
     'Mevcut Madde (Onaylandı)'),
    (15,'Audit Policies','Audit system event',
     'Success, Failure',
     'GPO: Advanced Audit Policy Configuration','Must',
     'Sistem kapanmaları, yeniden başlatmalar ve güvenlik alt sistemi olaylarının kaydı.',
     'Mevcut Madde (Onaylandı)'),
    (16,'Audit Policies','Audit account management',
     'Success, Failure',
     'GPO: Advanced Audit Policy Configuration','Must',
     'Hesap oluşturma, silme ve değişikliklerinin izlenmesi.',
     'Mevcut Madde (Onaylandı)'),
    (17,'Security Options','Additional restrictions for anonymous connections',
     '"Do not allow anonymous enumeration of SAM accounts and shares" — domain ve bağımsız sunucular için. RestrictAnonymous = 2.',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'Null session üzerinden AD ve paylaşım keşfini engeller.',
     'Mevcut Madde (Registry Değeri Eklendi)'),
    (18,'Security Options','Shutdown: Allow system to be shut down without having logged on',
     'Disabled — ShutdownWithoutLogon = 0',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'Yetkisiz fiziksel erişimde sunucunun kapatılmasını engeller.',
     'Mevcut Madde (Registry Değeri Eklendi)'),
    (19,'Security Options','Disconnect clients when logon hours expire',
     'Enabled — ForceLogoffWhenHourExpire = 1',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'Çalışma saati dışında oturum açık kalmasını engeller.',
     'Mevcut Madde (Registry Değeri Eklendi)'),
    (20,'Security Options','SMB Signing — Network client and server digitally sign communications (always)',
     'RequireSecuritySignature = 1 (hem client hem server için). EnableSecuritySignature = 1.',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'SMB MITM ve relay saldırılarını engeller. NTLM relay\'in temel önlemi.',
     'Mevcut Madde (Güçlendirildi: Her İki Taraf Zorunlu)'),
    (21,'Security Options','Interactive logon: Do not require CTRL+ALT+DEL',
     'Disabled — DisableCAD = 0 (CTRL+ALT+DEL zorunlu olsun)',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'Sahte giriş ekranı saldırılarını (logon spoofing) engeller.',
     'Mevcut Madde (Değer Netleştirildi)'),
    (22,'Security Options','LAN Manager authentication level',
     'Send NTLMv2 response only, refuse LM & NTLM — LmCompatibilityLevel = 5',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'Zayıf LM ve NTLM v1 kimlik doğrulama protokollerini devre dışı bırakır.',
     'Mevcut Madde (Registry Değeri Eklendi)'),
    (23,'Security Options','Domain member: Digitally encrypt or sign secure channel data',
     'Always: RequireSignOrSeal = 1 | When possible: SealSecureChannel = 1 | Sign: SignSecureChannel = 1',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'Netlogon kanalının şifreli ve imzalı olmasını zorunlu kılarak MITM saldırılarını engeller.',
     'Mevcut Madde (Güçlendirildi: 3 Ayar Birleştirildi)'),
    (24,'Security Options','Prevent users from installing printer drivers',
     'Enabled — AddPrinterDrivers = 1',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'Non-admin kullanıcıların yazıcı sürücüsü kurarak privilege escalation yapmasını engeller (PrintNightmare).',
     'Mevcut Madde (PrintNightmare Bağlamı Eklendi)'),
    (25,'System Service Security','Disable unnecessary and dangerous services.',
     'Kapatılacaklar: Print Spooler (sunucularda), Remote Registry, Bluetooth, Xbox servisleri, AllJoyn Router.\nKapatılması değerlendirilecekler: WinRM (kullanılmıyorsa), Telnet, TFTP.',
     'GPO: Computer Configuration > Windows Settings > Security Settings > System Services','Must',
     'Gereksiz servisler saldırı yüzeyi oluşturur. Print Spooler PrintNightmare, Remote Registry yetkisiz kayıt erişimi riskidir.',
     'GÜNCELLEME: Eski ve mevcut olmayan servisler kaldırıldı, güncel tehdit vektörleri eklendi.'),
    (26,'Security Hardening','Cmd and PowerShell logs transferred to SIEM.',
     'ProcessCreationIncludeCmdLine_Enabled = 1 (bkz. Madde 12). PowerShell ScriptBlockLogging ve ModuleLogging GPO ile etkinleştirilmeli.',
     'GPO: Administrative Templates > Windows Components > Windows PowerShell','Must',
     'Saldırganların PowerShell üzerinden yürüttüğü komutların tam kaydının SIEM\'e iletilmesi. Forensic ve olay müdahalesi için kritik.',
     'Mevcut Madde (PowerShell Detayı Eklendi)'),
    (27,'Security Settings','Solicited and Unsolicited Remote Assistance both disabled.',
     'fAllowToGetHelp = 0, fAllowUnsolicited = 0',
     'GPO: Computer Configuration > Administrative Templates > System > Remote Assistance','Must',
     'Uzaktan yardım özelliğinin yetkisiz erişim vektörü olarak kullanılmasını engeller.',
     'GÜNCELLEME: Sadece Solicited değil, Unsolicited da eklendi.'),
    (28,'Security Settings','Turn off Autoplay for all drives.',
     'NoAutorun = 1, NoAutoplayForNonVolume = 1',
     'GPO: Computer Configuration > Administrative Templates > Windows Components > AutoPlay Policies','Must',
     'USB ve diğer taşınabilir medyadan otomatik kod çalıştırılmasını (BadUSB) engeller.',
     'Mevcut Madde (Registry Değerleri Eklendi)'),
    (29,'Terminal Services (RDP)','RDP encryption level set to High and NLA enforced.',
     'MinEncryptionLevel = 3 (High). UserAuthentication = 1 (NLA zorunlu). fEncryptRPCTraffic = 1.',
     'GPO: Computer Configuration > Administrative Templates > Windows Components > Remote Desktop Services','Must',
     'GÜNCELLEME: "Client Compatible" yetersiz. NLA olmadan BlueKeep ve brute-force riski var. High Level zorunlu kılındı.',
     'GÜNCELLEME: Client Compatible -> High Level + NLA Eklendi'),
    (30,'Terminal Services (RDP)','Do not allow passwords to be saved in RDP.',
     'DisablePasswordSaving = 1',
     'GPO: Computer Configuration > Administrative Templates > Windows Components > Remote Desktop Services','Must',
     'Kaydedilmiş RDP şifrelerinin çalınarak lateral movement yapılmasını engeller.',
     'Mevcut Madde (Registry Değeri Eklendi)'),
    # Yeni eklenenler
    (31,'Security Hardening','WDigest authentication must be disabled to prevent cleartext credentials in memory.',
     'HKLM\\System\\CurrentControlSet\\Control\\SecurityProviders\\WDigest\nUseLogonCredential = 0',
     'GPO: Registry / Preference','Must',
     'WDigest açık olduğunda Windows şifreleri bellekte açık metin olarak saklar. Mimikatz ile saniyeler içinde çalınabilir.',
     'YENİ EKLENDİ (Kritik - CIS 18.3.7 / Mimikatz Savunması)'),
    (32,'Security Hardening','SMBv1 protocol must be completely disabled.',
     'Set-SmbServerConfiguration -EnableSMB1Protocol $false\nKayıt defteri: HKLM\\SYSTEM\\CurrentControlSet\\Services\\mrxsmb10 -> Start = 4',
     'PowerShell / Registry / GPO','Must',
     'SMBv1, 2017 WannaCry/EternalBlue saldırısının kullandığı protokoldür. Modern ortamda kullanım gerekçesi yoktur.',
     'YENİ EKLENDİ (Kritik - CIS 18.3.3 / EternalBlue Savunması)'),
    (33,'Security Hardening','UAC (User Account Control) must be enabled and configured securely.',
     'EnableLUA = 1\nConsentPromptBehaviorAdmin = 2 (onay gereksin)\nPromptOnSecureDesktop = 1',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'UAC kapalıysa her program sessizce admin yetkisi alır. Tüm privilege escalation zincirinin temel savunmasıdır.',
     'YENİ EKLENDİ (Kritik - CIS 18.9.97 UAC)'),
    (34,'Security Hardening','LSASS must run as Protected Process Light (PPL) to prevent credential dumping.',
     'HKLM\\SYSTEM\\CurrentControlSet\\Control\\Lsa\nRunAsPPL = 1',
     'GPO: Registry / Preference','Must',
     'LSASS PPL ile çalıştığında Mimikatz ve benzeri araçlar bellekten hash/şifre çekemez.',
     'YENİ EKLENDİ (Kritik - CIS 18.3.1 / Credential Dumping Savunması)'),
    (35,'Security Hardening','EventLog maximum size must be configured to prevent log overwrite.',
     'Security Log: MaxSize = 196608 KB (192 MB)\nApplication Log: MaxSize = 32768 KB\nSystem Log: MaxSize = 32768 KB',
     'GPO: Computer Configuration > Windows Settings > Security Settings > Event Log','Must',
     'Log boyutu küçük kalırsa eski olaylar silinir, saldırı izleri kaybolur. SIEM korelasyonu için yeterli geçmiş gereklidir.',
     'YENİ EKLENDİ (CIS 18.9.26 / Forensic Görünürlük)'),
    (36,'Security Hardening','NTLM minimum security level must be enforced for client and server.',
     'NtlmMinClientSec = 537395200\nNtlmMinServerSec = 537395200\n(NTLMv2 + 128-bit şifreleme zorunlu)',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'Zayıf NTLM şifrelemelere karşı downgrade saldırılarını engeller.',
     'YENİ EKLENDİ (CIS 2.3.11 / NTLM Downgrade Savunması)'),
    (37,'Security Hardening','Print Spooler service must be disabled on servers that do not require printing.',
     'Set-Service -Name Spooler -StartupType Disabled\nStop-Service -Name Spooler',
     'PowerShell / GPO: System Services','Must',
     'Print Spooler, PrintNightmare (CVE-2021-34527) güvenlik açığının kullandığı servistir. Yazıcı gerektirmeyen sunucularda kesinlikle kapatılmalıdır.',
     'YENİ EKLENDİ (Kritik - PrintNightmare / CVE-2021-34527)'),
    (38,'Security Hardening','UNC Hardened Paths must be configured for NETLOGON and SYSVOL.',
     '\\\\*\\NETLOGON: RequireMutualAuthentication=1, RequireIntegrity=1\n\\\\*\\SYSVOL: RequireMutualAuthentication=1, RequireIntegrity=1',
     'GPO: Administrative Templates > Network > Network Provider > Hardened UNC Paths','Must',
     'NETLOGON ve SYSVOL üzerinden GPO tampering ve MS15-011 açığına karşı koruma sağlar.',
     'YENİ EKLENDİ (Kritik - CIS 18.5.14.1 / GPO Tampering Savunması)'),
    (39,'Security Hardening','Kerberos encryption must be limited to AES only; RC4 must be disabled.',
     'SupportedEncryptionTypes = 2147483640\n(AES128 + AES256 + Future; RC4 hariç)',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'RC4 şifrelemesi Kerberoasting saldırılarında kullanılır. AES zorunlu kılınarak offline şifre kırma saldırısı zorlaştırılır.',
     'YENİ EKLENDİ (Kritik - CIS 18.3.9 / Kerberoasting Savunması)'),
    (40,'Security Hardening','Credential Guard / Virtualization Based Security (VBS) must be enabled where hardware supports.',
     'EnableVirtualizationBasedSecurity = 1\nRequirePlatformSecurityFeatures = 1\nLsaCfgFlags = 1 (Credential Guard)',
     'GPO: Computer Configuration > Administrative Templates > System > Device Guard','Must',
     'Credential Guard aktifken LSASS ayrı bir sanallaştırılmış ortamda çalışır, Mimikatz hash çekemez. TPM 2.0 ve Secure Boot gerektirir.',
     'YENİ EKLENDİ (Kritik - CIS 18.8.5 / Credential Guard)'),
]

# Section grupları - server
SRV_SECTIONS = {
    1:  'Installation',
    4:  'User Account & Password Policy',
    9:  'Audit Policies',
    17: 'Security Options',
    25: 'System Services & Hardening',
    31: 'New Security Controls (CIS Aligned)',
}

r = 3
for row_data in server_data:
    no = row_data[0]
    if no in SRV_SECTIONS:
        make_section_row(ws_srv, r, SRV_SECTIONS[no])
        r += 1
    is_new = 'YENİ EKLENDİ' in str(row_data[7])
    is_rev = 'GÜNCELLEME' in str(row_data[7]) or 'Güçlendirildi' in str(row_data[7]) or 'Değer' in str(row_data[7])
    write_row(ws_srv, r, list(row_data), is_new=is_new, is_revised=is_rev)
    r += 1

r += 1
ws_srv.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
c = ws_srv.cell(r, 1, '  Renk Kodu:   Beyaz/Gri = Mevcut madde   |   Sarı = Güncellendi / Güçlendirildi   |   Turuncu = YENİ EKLENDİ')
c.font = fnt(True, 8, '7F3F00')
c.fill = fill('FFF0CC')
c.alignment = aln('left', False, 'center')
ws_srv.row_dimensions[r].height = 18

# ══════════════════════════════════════════════════════════════════
# SHEET 3 — WINDOWS CLIENT / DESKTOP
# ══════════════════════════════════════════════════════════════════
ws_cli = wb.create_sheet('Windows Client Baseline')
setup_columns(ws_cli)
make_header_row(ws_cli, 1, 'Configuration Baseline Checklist — Windows Client / Desktop  |  CIS L1 Aligned')
make_col_headers(ws_cli, 2)

client_data = [
    (1,'Installation','Ensure the source of operating system media is legal.',
     'Kurulum medyası yalnızca Microsoft\'un resmi kaynağından veya kurumsal SCCM/MDT dağıtım altyapısından sağlanmalıdır.',
     'Microsoft Volume Licensing / MDT / SCCM','Must',
     'Trojanlanmış işletim sistemi kurulumlarını engeller.',
     'Mevcut Madde (Kaynak Netleştirildi)'),
    (2,'Installation','Install the OS in an independent, secure partition; use NTFS only.',
     'C: sürücüsü NTFS formatlanmalı. Sistem ve kullanıcı verileri tercihen ayrı bölümlerde tutulmalıdır. Şifrelenmemiş FAT/FAT32 bölümler olmamalıdır.',
     'Disk Management / MDT','Must',
     'FAT dosya sistemi NTFS güvenlik izinlerini desteklemez; veri izolasyonu sağlanmış olur.',
     'Mevcut Madde (NTFS Zorunluluğu Güçlendirildi)'),
    (3,'Installation','The name of clients must meet relevant naming rules.',
     'Bilgisayar adı kurumsal hostname kurallarına uymalıdır. Hostname üzerinden işletim sistemi bilgisi açığa çıkmamalıdır.',
     'Active Directory / MDT / SCCM','Must',
     'Tutarlı isimlendirme envanter yönetimi ve SIEM korelasyonu için gereklidir.',
     'Mevcut Madde (Onaylandı)'),
    (4,'User Account Security Control','Administrator account shall be renamed.',
     'NewAdministratorName GPO ayarı ile varsayılan "Administrator" hesabı farklı isimle değiştirilmeli.',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'Otomatik saldırıların hedeflediği varsayılan hesap adını geçersiz kılar.',
     'Mevcut Madde (GPO Yolu Eklendi)'),
    (5,'User Account Security Control','Guest account shall be disabled.',
     'Guest hesabı Disabled olmalı. Etkinleştirilmesi gereken özel durum yoksa tamamen kaldırılabilir.',
     'GPO: Computer Configuration > Windows Settings > Security Settings','Must',
     'Anonim erişim riskini sıfırlar.',
     'Mevcut Madde (Onaylandı)'),
    (6,'Account Security Control','No local accounts for clients (except built-in); all access via AD.',
     'Yönetim için oluşturulan lokal hesaplar dışında (LAPS yönetimli built-in admin) lokal hesap bulunmamalıdır. Tüm kullanıcı girişleri AD hesabıyla yapılmalıdır.',
     'Active Directory / Local User Management','Must',
     'Lokal hesaplar merkezi yönetim dışında kalır; şifre rotasyonu ve denetim yapılamaz.',
     'Mevcut Madde (Kapsam Netleştirildi)'),
    (7,'Account Security Control','All local administrator accounts\' passwords managed by LAPS.',
     'Microsoft LAPS kurulmalı. ADPasswordEncryptionEnabled = 1. PasswordLength minimum 15 karakter. PasswordAgeDays = 30.',
     'Microsoft LAPS / Active Directory','Must',
     'Tüm istemcilerde aynı lokal admin şifresinin kullanılmasını engeller; lateral movement riskini azaltır.',
     'YENİ EKLENDİ (CIS / LAPS Zorunluluğu)'),
    (8,'Password Policy','Enable password complexity and enforce history.',
     'Karmaşıklık: Etkin. Minimum uzunluk: 21 karakter. Maksimum yaş: 90 gün. Geçmiş: 24 şifre. Kilitleme: 5 deneme / 15 dk.',
     'GPO: Default Domain Policy > Password Policy & Account Lockout Policy','Must',
     'Zayıf şifre kullanımını ve tekrar kullanımını engeller.',
     'Mevcut Madde (Tüm Değerler Birleştirildi)'),
    (9,'Account Lockout Policy','Account lockout threshold: 5 attempts, lockout duration: 15 minutes.',
     'LockoutBadCount = 5\nLockoutDuration = 15 (dakika)\nResetLockoutCount = 15 (dakika)',
     'GPO: Default Domain Policy > Account Lockout Policy','Must',
     'Brute-force şifre denemelerini otomatik kilitleyerek durdurur.',
     'Mevcut Madde (Değerler Güncellendi: 30 dk -> 15 dk, CIS uyumlu)'),
    (10,'Audit Policies','Audit logon events, account management, privilege use and process creation.',
     'Logon: Success/Failure | Account Management: Success/Failure | Privilege Use: Success/Failure | Process Creation: Success + CmdLine enabled',
     'GPO: Advanced Audit Policy Configuration','Must',
     'İstemcilerdeki anormal davranışların SIEM üzerinden tespiti için minimum log gereksinimi.',
     'Mevcut Madde (Process Creation Eklendi, Optional -> Must)'),
    (11,'User Rights Assignment','Access this computer from network — restricted.',
     'Administrators ve Authenticated Users. Power Users, Backup Operators ve Guest kaldırılmalıdır.',
     'GPO: User Rights Assignment','Must',
     'Gereksiz hesapların ağ üzerinden erişimini kısıtlar.',
     'Mevcut Madde (Onaylandı)'),
    (12,'User Rights Assignment','Add workstation to domain — Domain Admins only (NOT Authenticated Users).',
     'SeMachineAccountPrivilege yalnızca Domain Admins veya belirlenmiş bir IT grubu için tanımlanmalıdır. Authenticated Users kaldırılmalıdır.',
     'GPO: User Rights Assignment','Must',
     'GÜNCELLEME: Authenticated Users domain\'e makine ekleyebilirse saldırgan da ekleyebilir (MachineAccountQuota saldırısı). Sadece admin yapabilmeli.',
     'GÜNCELLEME: Authenticated Users -> Domain Admins Only (Hatalı madde düzeltildi)'),
    (13,'Security Options','SMB Signing — Network client and server must always digitally sign communications.',
     'RequireSecuritySignature = 1 | EnableSecuritySignature = 1 (hem client hem server)',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'SMB relay saldırılarını engeller. NTLM relay\'in en temel önlemi.',
     'Mevcut Madde (Güçlendirildi: Her İki Yön Zorunlu)'),
    (14,'Security Options','LAN Manager: Send NTLMv2 only, refuse LM & NTLM.',
     'LmCompatibilityLevel = 5',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'Zayıf LM ve NTLMv1 protokollerini devre dışı bırakır.',
     'YENİ EKLENDİ (Sunucu listesinde vardı, client listesinde eksikti)'),
    (15,'Security Options','Interactive logon: Do not display last username.',
     'DontDisplayLastUserName = 1',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'Oturum açma ekranında önceki kullanıcı adının görünmesini engeller; saldırgan geçerli kullanıcı adını öğrenemez.',
     'YENİ EKLENDİ (CIS 2.3.7.4)'),
    (16,'Security Hardening','Block remote access to registry.',
     'Remote Registry servisi kapatılmalı (Disabled). Gerekli ise kısıtlı IP\'lerden erişime izin verilmeli.',
     'GPO: System Services / Registry ACL','Must',
     'Kayıt defterine uzaktan yetkisiz erişimi engeller.',
     'Mevcut Madde (Onaylandı)'),
    (17,'Security Hardening','Restrict null session access to named pipes and shares.',
     'RestrictNullSessAccess = 1\nNullSessionPipes = (boş)\nNullSessionShares = (boş)',
     'GPO: Security Settings > Local Policies > Security Options','Must',
     'Kimliksiz bağlantılarla ağ paylaşımı ve pipe erişimini engeller.',
     'Mevcut Madde (Registry Değerleri Eklendi)'),
    (18,'Security Hardening','Turn off Autoplay and Autorun for all drives.',
     'NoAutorun = 1\nNoAutoplayForNonVolume = 1',
     'GPO: Administrative Templates > Windows Components > AutoPlay Policies','Must',
     'USB BadUSB ve autorun saldırılarını engeller.',
     'Mevcut Madde (Registry Değerleri Eklendi)'),
    (19,'Security Hardening','Cmd and PowerShell access restricted for end users; logs forwarded to SIEM.',
     'End user OU\'sunda: cmd.exe ve powershell.exe AppLocker veya SRP ile kısıtlanmalı.\nPowerShell ScriptBlockLogging ve ModuleLogging GPO ile aktif edilmeli.',
     'AppLocker / GPO: Administrative Templates > Windows Components > Windows PowerShell','Must',
     'Son kullanıcının komut satırı üzerinden zararlı komut çalıştırmasını engeller; PowerShell logları saldırı analizinde kritik.',
     'Mevcut Madde (2021 tarihi kaldırıldı, AppLocker eklendi)'),
    (20,'Security Hardening','Application Whitelist enforced via AppLocker or Windows Defender Application Control.',
     'AppLocker Publisher/Hash kuralları ile sadece onaylı uygulamaların çalışmasına izin verilmeli. Varsayılan: tüm bilinmeyenler bloklanmalı.',
     'GPO: Computer Configuration > Windows Settings > Security Settings > Application Control Policies','Must',
     'Yetkisiz ve zararlı yazılımların çalıştırılmasını engeller. Ransomware savunmasında en etkili önlemlerden biridir.',
     'Mevcut Madde (Detay ve Araç Eklendi)'),
    (21,'Security Hardening','WDigest must be disabled; LSASS must run as PPL.',
     'UseLogonCredential = 0 (WDigest)\nRunAsPPL = 1 (LSASS Protected Process)',
     'GPO: Registry Preference','Must',
     'Mimikatz ve benzeri araçların bellekten şifre/hash çalmasını engeller.',
     'YENİ EKLENDİ (Kritik - Credential Theft Savunması)'),
    (22,'Security Hardening','BitLocker full disk encryption must be enabled on all client devices.',
     'BitLocker: Enabled\nEncryption Method: XTS-AES 256\nTPM + PIN veya TPM + Startup Key\nRecovery Key AD\'de saklanmalıdır.',
     'GPO: Computer Configuration > Administrative Templates > Windows Components > BitLocker Drive Encryption','Must',
     'Cihaz çalındığında veya kaybolduğunda disk içeriğinin okunmasını engeller. Laptop/mobil cihazlar için kritik.',
     'YENİ EKLENDİ (Kritik - CIS 18.9.11 / Kayıp Cihaz Savunması)'),
    (23,'Security Hardening','Windows Defender SmartScreen must be enabled.',
     'EnableSmartScreen = 1\nShellSmartScreenLevel = Block\nPreventOverride = 1 (kullanıcı geçemesin)',
     'GPO: Administrative Templates > Windows Components > Windows Defender SmartScreen','Must',
     'Bilinmeyen ve şüpheli dosya/URL\'lerin çalıştırılmasını engeller; phishing koruması sağlar.',
     'YENİ EKLENDİ (CIS 18.9.85 / SmartScreen)'),
    (24,'Security Hardening','Copilot and consumer AI features must be disabled on corporate devices.',
     'TurnOffWindowsCopilot = 1',
     'GPO: Administrative Templates > Windows Components > Windows Copilot','Must',
     'Kurumsal verinin Microsoft bulut AI hizmetlerine gönderilmesini engeller. KVKK ve veri gizliliği uyumu için gereklidir.',
     'YENİ EKLENDİ (Win11 / Veri Gizliliği)'),
    (25,'Security Hardening','Remote Desktop (RDP) must use NLA and High encryption; password saving disabled.',
     'UserAuthentication = 1 (NLA)\nMinEncryptionLevel = 3 (High)\nDisablePasswordSaving = 1',
     'GPO: Administrative Templates > Windows Components > Remote Desktop Services','Must',
     'NLA olmadan BlueKeep açığı ve brute-force riski oluşur. Kaydedilmiş şifreler lateral movement için kullanılabilir.',
     'YENİ EKLENDİ (Client RDP Güvenliği - CIS)'),
]

CLI_SECTIONS = {
    1:  'Installation',
    4:  'User Account Security Control',
    8:  'Password & Account Lockout Policy',
    10: 'Audit Policies',
    11: 'User Rights Assignment',
    13: 'Security Options',
    16: 'Security Hardening',
    21: 'New Security Controls (CIS Aligned)',
}

r = 3
for row_data in client_data:
    no = row_data[0]
    if no in CLI_SECTIONS:
        make_section_row(ws_cli, r, CLI_SECTIONS[no])
        r += 1
    is_new = 'YENİ EKLENDİ' in str(row_data[7])
    is_rev = 'GÜNCELLEME' in str(row_data[7]) or 'Güçlendirildi' in str(row_data[7]) or 'Değer' in str(row_data[7])
    write_row(ws_cli, r, list(row_data), is_new=is_new, is_revised=is_rev)
    r += 1

r += 1
ws_cli.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
c = ws_cli.cell(r, 1, '  Renk Kodu:   Beyaz/Gri = Mevcut madde   |   Sarı = Güncellendi / Güçlendirildi   |   Turuncu = YENİ EKLENDİ')
c.font = fnt(True, 8, '7F3F00')
c.fill = fill('FFF0CC')
c.alignment = aln('left', False, 'center')
ws_cli.row_dimensions[r].height = 18

# Sheet sırası
wb._sheets = [ws_linux, ws_srv, ws_cli]

script_dir = os.path.dirname(os.path.abspath(__file__))
out = os.path.join(script_dir, 'Configuration_Baseline_Checklist.xlsx')
wb.save(out)
print('OK:', out)
print(f'Linux: {len(linux_data)} madde | Server: {len(server_data)} madde | Client: {len(client_data)} madde')
